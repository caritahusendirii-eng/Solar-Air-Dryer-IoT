from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from datetime import date

OUT = r"C:\Users\Abel Rizky\OneDrive\Documents\Tugas Akhir\aplikasi\smart-ac-dashboard\Project_Review_Solar_Dryer_IoT.docx"

NAVY = "17324D"
BLUE = "276FBF"
TEAL = "0F766E"
GOLD = "D97706"
RED = "B42318"
PALE = "EAF2F8"
LIGHT = "F5F7FA"
GRAY = "52606D"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def cell_margin(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); mar = tcPr.first_child_found_in('w:tcMar')
    if mar is None: mar = OxmlElement('w:tcMar'); tcPr.append(mar)
    for side, value in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = mar.find(qn(f'w:{side}'))
        if node is None: node = OxmlElement(f'w:{side}'); mar.append(node)
        node.set(qn('w:w'), str(value)); node.set(qn('w:type'), 'dxa')

def set_cell_text(cell, text, bold=False, color="000000", size=9.5):
    cell.text = ''
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.space_before = Pt(0)
    r = p.add_run(str(text)); r.bold = bold; r.font.name = 'Aptos'; r._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); r._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos'); r.font.size = Pt(size); r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; cell_margin(cell)

def set_widths(table, widths):
    table.autofit = False
    for row in table.rows:
        for cell, width in zip(row.cells, widths): cell.width = Inches(width)

def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers)); table.alignment = WD_TABLE_ALIGNMENT.LEFT; table.style = 'Table Grid'
    set_widths(table, widths)
    for c,h in zip(table.rows[0].cells, headers):
        shade(c, NAVY); set_cell_text(c,h,True,'FFFFFF',9)
    for i,row in enumerate(rows):
        cells = table.add_row().cells
        for c,v in zip(cells,row):
            if i % 2 == 1: shade(c, LIGHT)
            set_cell_text(c,v,False,'1F2937',9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table

def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.add_run(item)

def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f'Heading {level}')
    p.add_run(text)
    return p

def add_callout(doc, label, text, fill='EAF2F8', label_color=NAVY):
    t=doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.LEFT; set_widths(t,[6.5]); c=t.cell(0,0); shade(c,fill); cell_margin(c,160,180,160,180)
    c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(0)
    r=p.add_run(label+'  '); r.bold=True; r.font.color.rgb=RGBColor.from_string(label_color); r.font.size=Pt(10)
    r=p.add_run(text); r.font.size=Pt(10); r.font.color.rgb=RGBColor.from_string('243B53')
    doc.add_paragraph().paragraph_format.space_after=Pt(1)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(0.78); sec.bottom_margin=Inches(0.72); sec.left_margin=Inches(0.8); sec.right_margin=Inches(0.7)
sec.header_distance=Inches(0.3); sec.footer_distance=Inches(0.35)

styles=doc.styles
normal=styles['Normal']; normal.font.name='Aptos'; normal._element.rPr.rFonts.set(qn('w:ascii'),'Aptos'); normal._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos'); normal.font.size=Pt(10); normal.font.color.rgb=RGBColor.from_string('243B53'); normal.paragraph_format.space_after=Pt(6); normal.paragraph_format.line_spacing=1.12
for level,size,color in [(1,16,NAVY),(2,12,BLUE),(3,10.5,TEAL)]:
    s=styles[f'Heading {level}']; s.font.name='Aptos Display'; s._element.rPr.rFonts.set(qn('w:ascii'),'Aptos Display'); s._element.rPr.rFonts.set(qn('w:hAnsi'),'Aptos Display'); s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color); s.paragraph_format.space_before=Pt(14 if level==1 else 9); s.paragraph_format.space_after=Pt(5)

# Header/footer
header=sec.header.paragraphs[0]; header.text='PROJECT REVIEW  |  SOLAR DRYER IoT'; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
for r in header.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('Dokumen review internal • ' + date.today().strftime('%d %B %Y'))
for r in footer.runs: r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(GRAY)

# Cover masthead
p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(30); p.paragraph_format.space_after=Pt(4)
r=p.add_run('PROJECT REVIEW DOCUMENT'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(11); r.font.color.rgb=RGBColor.from_string(TEAL)
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6)
r=p.add_run('Solar Dryer IoT'); r.bold=True; r.font.name='Aptos Display'; r.font.size=Pt(30); r.font.color.rgb=RGBColor.from_string(NAVY)
p=doc.add_paragraph('Tinjauan implementasi sistem monitoring dan pengendalian pengering surya berbasis IoT.'); p.paragraph_format.space_after=Pt(20)
for r in p.runs: r.font.size=Pt(13); r.font.color.rgb=RGBColor.from_string(GRAY)
add_table(doc,['METADATA','RINGKASAN'],[
    ['Objek review','Repositori smart-ac-dashboard (implementasi Solar Dryer IoT)'],
    ['Komponen','ESP32, Firebase Realtime Database, React/Capacitor, Flutter'],
    ['Tanggal review',date.today().strftime('%d %B %Y')],
    ['Metode','Review statis dokumentasi, konfigurasi, dan kode sumber; tanpa uji perangkat fisik'],
], [1.55,4.95])
add_callout(doc,'KESIMPULAN EKSEKUTIF','Sistem telah memiliki fondasi fungsional yang kuat untuk prototipe/skripsi: pembacaan multi-sensor, kontrol PWM otomatis/manual, histori data, serta dashboard web dan mobile. Kesiapan produksi masih tertahan oleh temuan keamanan akses Firebase dan pengelolaan rahasia yang perlu diselesaikan lebih dahulu.','E7F4F1',TEAL)
add_heading(doc,'Ringkasan Penilaian',1)
add_table(doc,['ASPEK','STATUS','CATATAN'],[
 ['Fungsi inti','Baik','Alur sensor → cloud → dashboard → aktuator tercakup.'],
 ['Arsitektur','Baik','Pemisahan firmware, web, mobile, dan backend jelas.'],
 ['Keamanan','Perlu tindakan','Aturan database terlalu terbuka dan rahasia ditemukan di source.'],
 ['Keandalan','Cukup','Ada reconnect Wi-Fi, NTP/RTC, dan logging SD; belum tampak strategi retry/monitoring cloud yang menyeluruh.'],
 ['Kualitas rilis','Cukup','Tersedia APK dan tangkapan layar; automated test/CI belum tampak.'],
], [1.25,1.2,4.05])

doc.add_page_break()
add_heading(doc,'1. Ruang Lingkup dan Metode Review',1)
doc.add_paragraph('Review ini memeriksa kode dan dokumentasi yang tersedia pada repositori, terutama firmware ESP32, aplikasi React, aplikasi Flutter, konfigurasi Firebase, serta dokumen arsitektur dan fungsi fitur. Penilaian tidak mencakup pengukuran sensor di lapangan, pengujian beban, audit Firebase Console, atau pentest dinamis.')
add_heading(doc,'Artefak utama yang ditinjau',2)
add_table(doc,['AREA','ARTEFAK','PERAN'],[
 ['Firmware','esp32/smart_ac.ino','Membaca 10 sensor, mengendalikan dua blower dengan PWM, menulis data/histori ke Firebase dan SD card.'],
 ['Web','react-dashboard/src','Dashboard React dengan autentikasi, peran pengguna, monitoring, kontrol, ekspor, dan manajemen.'],
 ['Mobile','flutter_dashboard/lib','Dashboard Flutter untuk monitoring, kontrol blower, dan visualisasi histori.'],
 ['Backend','database.rules.json','Aturan Firebase Realtime Database dan pembatasan beberapa jalur data.'],
 ['Dokumentasi','README.md dan docs/','Arsitektur, fitur, langkah penggunaan, serta lampiran kode.'],
], [1.0,2.05,3.45])
add_heading(doc,'2. Gambaran Arsitektur',1)
doc.add_paragraph('Arsitektur menggunakan Firebase Realtime Database sebagai penghubung real-time. ESP32 membaca sensor DHT22 dan DS18B20, menghitung nilai rata-rata, menjalankan aturan kontrol otomatis, lalu menyimpan data saat ini dan histori. Aplikasi React/Capacitor dan Flutter membaca data yang sama; perintah mode dan kecepatan blower ditulis kembali ke database untuk diproses firmware.')
add_table(doc,['LAPISAN','KOMPONEN','TANGGUNG JAWAB'],[
 ['Perangkat','ESP32, DHT22, DS18B20, BTS7960, DS3231, SD card','Akuisisi suhu/kelembapan, penjadwalan lokal, PWM blower, cache log.'],
 ['Cloud','Firebase RTDB + Auth','Sinkronisasi data/perintah real-time dan identitas pengguna.'],
 ['Web','React 18 + Recharts + Capacitor','Dashboard operasi, kontrol master, histori/ekspor, manajemen peran.'],
 ['Mobile','Flutter + Firebase Database + fl_chart','Monitoring dan kontrol dari perangkat Android.'],
], [1.1,2.1,3.3])
add_callout(doc,'ALUR DATA','Sensor → ESP32 → /sensor_data dan /history → Web/Mobile. Perintah pengguna → /control_mode dan /blowers → stream Firebase pada ESP32 → PWM driver blower.','EAF2F8',NAVY)

add_heading(doc,'3. Cakupan Fungsional yang Teridentifikasi',1)
add_bullets(doc,[
 'Monitoring real-time suhu dan kelembapan, termasuk indikator konektivitas dan pembaruan terakhir.',
 'Agregasi hingga 10 sensor: lima DHT22 (suhu/kelembapan) dan lima DS18B20 (suhu).',
 'Mode otomatis dengan tingkat PWM blower: 0%, 25%, 50%, 75%, dan 100% berdasarkan ambang suhu rata-rata.',
 'Mode manual dengan pengaturan kecepatan blower; web menyediakan kontrol master untuk menyelaraskan dua blower.',
 'Riwayat data di Firebase, pencatatan CSV pada SD card, grafik tren, serta ekspor data di dashboard web.',
 'Autentikasi email/password, peran akses, status kehadiran, dan layar manajemen pengguna pada dashboard web.',
 'Mode pemeliharaan dan overload serta halaman operasional terkait pada dashboard web.',
])

doc.add_page_break()
add_heading(doc,'4. Temuan Positif',1)
add_table(doc,['ID','TEMUAN','DAMPAK'],[
 ['S-01','Pemisahan komponen perangkat, backend, web, dan mobile cukup jelas.','Memudahkan pengembangan dan penelusuran masalah.'],
 ['S-02','Firmware memiliki mekanisme reconnect Wi-Fi, sinkronisasi NTP ke RTC, dan fallback logging SD card.','Meningkatkan ketahanan saat koneksi cloud tidak stabil.'],
 ['S-03','Kontrol otomatis memakai ambang bertingkat dan UI menjelaskan ambang tersebut.','Perilaku operasional lebih dapat dipahami pengguna.'],
 ['S-04','Dashboard web menyediakan autentikasi, model peran, pemisahan tab, dan manajemen akses.','Fondasi kontrol akses sudah tersedia di aplikasi.'],
 ['S-05','Dokumentasi menyertakan arsitektur, komponen, dan petunjuk penggunaan; tersedia artefak APK/screenshot.','Mendukung demonstrasi dan serah terima proyek.'],
], [0.65,3.35,2.5])
add_heading(doc,'5. Temuan Risiko dan Rekomendasi',1)
add_table(doc,['PRIORITAS','TEMUAN','REKOMENDASI'],[
 ['P0 – Kritis','Kredensial Wi-Fi dan akun Firebase tertanam langsung di firmware yang tersimpan dalam repositori.','Segera rotasi semua kredensial yang sudah terekspos; pindahkan konfigurasi ke file lokal yang diabaikan Git atau mekanisme provisioning; buat template konfigurasi tanpa nilai rahasia.'],
 ['P0 – Kritis','Rules Firebase memberi akses baca/tulis publik pada sensor_data, blowers, control_mode, dan history. Ini membuka kontrol aktuator serta perubahan/penghapusan data tanpa autentikasi.','Ubah rules menjadi default-deny; wajibkan auth untuk semua jalur sensitif; batasi write per peran; validasikan tipe dan rentang speed/mode di rules atau Cloud Functions.'],
 ['P1 – Tinggi','Reset histori memakai kata sandi statis di sisi klien dan operasi hapus dilakukan langsung ke /history.','Hapus password hard-coded; izinkan hanya admin melalui Firebase rules/Cloud Function; tambahkan konfirmasi berlapis dan audit trail.'],
 ['P1 – Tinggi','Flutter membaca skema lama (temperature, humidity, timestamp), sedangkan firmware aktif menulis avg_t, avg_h, tanggal/waktu, dan sensor_data.','Samakan kontrak data lintas klien; buat model data/versioning; tambahkan test integrasi untuk skema Firebase.'],
 ['P2 – Sedang','Perintah auto menulis speed kedua blower, namun firmware pada jalur auto hanya memanggil pengendalian fisik untuk indeks blower pertama.','Tegaskan kebutuhan operasional lalu panggil setBlowerSpeed untuk kedua blower atau ubah UI agar mencerminkan perilaku aktual.'],
 ['P2 – Sedang','Tidak tampak pipeline CI, test unit bermakna, atau validasi hardware-in-the-loop.','Tambahkan lint/build/test otomatis untuk React dan Flutter; uji fungsi ambang/PWM dalam unit test dan siapkan checklist uji lapangan.'],
], [0.85,2.55,3.1])
add_callout(doc,'CATATAN KEAMANAN','Nilai rahasia tidak disalin ke dokumen ini. Karena kredensial terlihat pada source, perlakukan sebagai telah terpapar dan lakukan rotasi sebelum demonstrasi atau publikasi repositori.','FDECEC',RED)

doc.add_page_break()
add_heading(doc,'6. Prioritas Perbaikan',1)
add_table(doc,['FASE','TARGET','AKSI UTAMA','HASIL YANG DIHARAPKAN'],[
 ['0–2 hari','Amankan akses','Rotasi credential; bersihkan riwayat Git bila akan dipublikasikan; terapkan rules default-deny.','Tidak ada akses publik ke perintah/riwayat dan tidak ada rahasia aktif di source.'],
 ['1 minggu','Stabilkan kontrak data','Definisikan skema Firebase tunggal; validasi batas speed 0–100 dan mode auto/manual.','Web, mobile, dan firmware membaca/menulis struktur yang konsisten.'],
 ['2 minggu','Tingkatkan keselamatan operasi','Tambahkan role enforcement server-side, audit log perintah, serta fail-safe saat data sensor invalid/koneksi hilang.','Kontrol aktuator dapat ditelusuri dan lebih aman.'],
 ['Berikutnya','Kualitas rilis','Tambah CI, test unit, panduan deployment, dan checklist uji perangkat.','Rilis lebih mudah direproduksi dan diverifikasi.'],
], [0.8,1.2,2.3,2.2])
add_heading(doc,'7. Kriteria Penerimaan Rilis Berikutnya',1)
add_bullets(doc,[
 'Semua kredensial aktif sudah dirotasi dan tidak tersimpan dalam source, APK, atau dokumentasi publik.',
 'Tidak ada jalur Firebase sensitif yang dapat ditulis anonim; kontrol blower memerlukan pengguna dan peran yang sah.',
 'Kontrak data sensor, blower, histori, dan status terdokumentasi serta digunakan konsisten oleh ESP32, React, dan Flutter.',
 'Pengujian memverifikasi batas PWM, mode otomatis/manual, sensor gagal baca, koneksi putus, dan pemulihan koneksi.',
 'Riwayat dan perubahan perintah dapat diaudit; penghapusan data dibatasi untuk peran yang tepat.',
])
add_heading(doc,'8. Kesimpulan',1)
doc.add_paragraph('Solar Dryer IoT menunjukkan implementasi yang cukup matang untuk tujuan demonstrasi dan penelitian terapan: fitur monitoring, histori, kontrol otomatis/manual, dan multi-platform telah hadir. Nilai terkuatnya adalah integrasi end-to-end dari sensor ke antarmuka pengguna. Namun, akses data dan kontrol yang masih terlalu terbuka membuat proyek belum layak dianggap siap produksi. Fokus pertama yang disarankan adalah menutup celah keamanan dan menyamakan kontrak data, kemudian melengkapi pengujian serta audit operasional.')
add_callout(doc,'STATUS REKOMENDASI','Layak untuk prototipe/demonstrasi setelah kredensial yang terekspos dirotasi. Belum direkomendasikan untuk deployment operasional/produksi sampai kontrol akses Firebase, pengelolaan rahasia, dan validasi lintas klien diperbaiki.','FFF4E5',GOLD)

doc.save(OUT)
print(OUT)
